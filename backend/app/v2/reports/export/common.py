from __future__ import annotations

"""Shared helpers for Word/Excel report export (labels, colors, charts, data quality)."""

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from app.v2.config import V2Settings

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VM_ALERT_FILL = "F4CCCC"
VM_ALERT_RATIO = 0.2
VM_ALERT_BYTES = 100 * 1024**3
DOCX_FONT_ASCII = "Noto Serif"
DOCX_FONT_EAST_ASIA = "Noto Serif CJK SC"
CHART_FONT_FAMILY = "Noto Serif CJK JP"
CHART_FONT_PATH = "/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc"
REPORT_PRODUCT_NAME = "存储容量预测平台"
REPORT_COVER_TITLE = "SMARTX超融合存储容量分析报告"
REPORT_COVER_SUBTITLE = "SMARTX HCI Storage Capacity Analysis Report"
XLSX_TEMPLATE_PATH = Path(__file__).parent.parent / "templates" / "customer_report.xlsx"
XLSX_FONT_NAME = "Noto Sans CJK SC"
XLSX_CLUSTER_TEMPLATE_SHEET = "__CLUSTER_TEMPLATE__"
ACCENT = "1A3C6E"
ACCENT_DARK = "1A3C6E"
ACCENT_LIGHT = "F0F4FA"
ACCENT_SOFT = "F7F9FC"
BORDER = "D9E2EF"
SUCCESS = "EAF7EF"
WARNING = "FFF4E5"
DANGER = "FDEAEA"
TEXT_DARK = "333333"
TEXT_MUTED = "999999"
GROWTH_BLUE = "007ACC"
RATIO_ORANGE = "E67E22"
RATIO_RED = "DC3535"
EMPTY_MONTH_VM_TEXT = "当前样本跨度区间暂无 VM 增长数据"
PERIODS = [("month", "较上月", 30), ("quarter", "较上季度", 90), ("year", "较上一年", 365)]


@dataclass(frozen=True)
class ReportPeriodProfile:
    days: int
    display_label: str
    window_growth_label: str
    vm_growth_title: str
    vm_empty_text: str
    short_advice_title: str
    focus_note: str


REPORT_PERIOD_PROFILES: dict[int, ReportPeriodProfile] = {
    7: ReportPeriodProfile(
        days=7,
        display_label="近 7 天",
        window_growth_label="近 7 天样本增长",
        vm_growth_title="短期突增 VM",
        vm_empty_text="当前 7 天窗口暂无明显 VM 增长数据",
        short_advice_title="短期（本周）",
        focus_note="本报告侧重识别短期突增、本日新建 VM 和日增长来源。",
    ),
    14: ReportPeriodProfile(
        days=14,
        display_label="近 14 天",
        window_growth_label="近 14 天样本增长",
        vm_growth_title="近两周增长 VM",
        vm_empty_text="当前 14 天窗口暂无明显 VM 增长数据",
        short_advice_title="短期（两周内）",
        focus_note="本报告侧重观察近两周增长变化和短期异常来源。",
    ),
    30: ReportPeriodProfile(
        days=30,
        display_label="近 30 天",
        window_growth_label="近 30 天样本增长",
        vm_growth_title="月度增长 VM",
        vm_empty_text="当前 30 天窗口暂无明显 VM 增长数据",
        short_advice_title="短期（本月）",
        focus_note="本报告侧重月度运营窗口、容量增长和 VM TOP 变化。",
    ),
    90: ReportPeriodProfile(
        days=90,
        display_label="近 90 天",
        window_growth_label="近 90 天样本增长",
        vm_growth_title="季度增长 VM",
        vm_empty_text="当前 90 天窗口暂无明显 VM 增长数据",
        short_advice_title="短期（本季度）",
        focus_note="本报告侧重季度趋势、预测可信度和容量风险。",
    ),
    180: ReportPeriodProfile(
        days=180,
        display_label="近 180 天",
        window_growth_label="近 180 天样本增长",
        vm_growth_title="中长期增长 VM",
        vm_empty_text="当前 180 天窗口暂无明显 VM 增长数据",
        short_advice_title="短期（半年内）",
        focus_note="本报告侧重中长期增长趋势和容量治理。",
    ),
    365: ReportPeriodProfile(
        days=365,
        display_label="近 365 天",
        window_growth_label="近 365 天样本增长",
        vm_growth_title="年度增长 VM",
        vm_empty_text="当前 365 天窗口暂无明显 VM 增长数据",
        short_advice_title="短期（年度巡检）",
        focus_note="本报告侧重年度容量规划；采集历史不足一年时按可用样本窗口计算。",
    ),
}



def report_period_profile(period_days: int | None) -> ReportPeriodProfile:
    try:
        days = int(period_days or 30)
    except (TypeError, ValueError):
        days = 30
    return REPORT_PERIOD_PROFILES.get(days, REPORT_PERIOD_PROFILES[30])



def _export_context(report: dict[str, Any], settings: V2Settings, period_days: int, extension: str) -> dict[str, str]:
    now = _local_now(settings)
    profile = report_period_profile(period_days)
    scope = report.get("scope") or {}
    clusters = report.get("clusters") or []
    if scope.get("cluster_id"):
        scope_label = _first_cluster_name(clusters) or str(scope["cluster_id"])
    elif scope.get("tower_id") is not None:
        scope_label = f"tower-{scope['tower_id']}"
    else:
        scope_label = "all"
    scope_slug = _slug(scope_label)
    filename = f"storage-forecast-{scope_slug}-{now.strftime('%Y%m%d-%H%M%S')}-{period_days}d.{extension}"
    return {
        "filename": filename,
        "scope_label": scope_label,
        "generated_at": now.strftime("%Y-%m-%d %H:%M:%S %Z"),
        "clusters": clusters,
        "report": report,
        "period_window": report.get("period_window") or {},
        "app_version": settings.app_version,
        "profile": profile,
    }



def _persist_report(content: bytes, settings: V2Settings, filename: str) -> tuple[bytes, str, Path, str]:
    settings.reports_dir.mkdir(parents=True, exist_ok=True)
    path = settings.reports_dir / Path(filename).name
    path.write_bytes(content)
    return content, filename, path, f"/api/admin/exports/reports/{quote(path.name)}"



def _add_cluster_growth_chart(document: Document, clusters: list[dict[str, Any]]) -> None:
    trend_chart = _scope_trend_line_chart(clusters, "容量使用率趋势")
    if trend_chart is not None:
        _add_figure(document, trend_chart, "图 1：容量使用趋势", width=6.4)
    else:
        document.add_paragraph("暂无足够历史数据生成容量趋势图。")
    top_chart = _cluster_top_growth_bar_chart(clusters, "Top 5 集群月增长量")
    if top_chart is not None:
        _add_figure(document, top_chart, "图 2：Top 5 集群月增长量", width=6.5)



def _add_figure(document: Document, image: BytesIO, caption: str, width: float) -> None:
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(image, width=Inches(width))



def _period_window_label(report: dict[str, Any]) -> str:
    window = _effective_report_window(report)
    start = str(window.get("start_at") or "")
    end = str(window.get("end_at") or "")
    if start and end:
        timezone = _report_timezone(report)
        parsed_start = _parse_report_datetime(start)
        parsed_end = _parse_report_datetime(end)
        if parsed_start and parsed_end:
            return f"{parsed_start.astimezone(timezone).date().isoformat()} - {parsed_end.astimezone(timezone).date().isoformat()}"
        return f"{start[:10]} - {end[:10]}"
    return f"近 {report.get('window_days', 30)} 天"



def _requested_report_window_label(report: dict[str, Any]) -> str:
    period = report.get("period_window") or {}
    days = period.get("days") or report.get("window_days") or 30
    start = str(period.get("start_at") or "")
    end = str(period.get("end_at") or "")
    try:
        days_value = int(days)
    except (TypeError, ValueError):
        days_value = 30
    if start and end:
        timezone = _report_timezone(report)
        parsed_start = _parse_report_datetime(start)
        parsed_end = _parse_report_datetime(end)
        if parsed_start and parsed_end:
            return (
                f"近 {days_value} 天（"
                f"{parsed_start.astimezone(timezone).date().isoformat()} - "
                f"{parsed_end.astimezone(timezone).date().isoformat()}）"
            )
        return f"近 {days_value} 天（{start[:10]} - {end[:10]}）"
    return f"近 {days_value} 天"



def _effective_report_window(report: dict[str, Any]) -> dict[str, Any]:
    timezone = _report_timezone(report)
    period = report.get("period_window") or {}
    data = report.get("data_window") or {}
    period_start = _parse_report_datetime(str(period.get("start_at") or ""))
    period_end = _parse_report_datetime(str(period.get("end_at") or ""))
    data_start = _parse_report_datetime(str(data.get("start_at") or ""))
    data_end = _parse_report_datetime(str(data.get("end_at") or ""))
    start_candidates = [value for value in [period_start, data_start] if value is not None]
    end_candidates = [value for value in [period_end, data_end] if value is not None]
    if start_candidates and end_candidates:
        start = max(start_candidates)
        end = min(end_candidates)
        if start <= end:
            return {"start_at": start.astimezone(timezone).isoformat(), "end_at": end.astimezone(timezone).isoformat()}
    if data_start and data_end:
        return {"start_at": data_start.astimezone(timezone).isoformat(), "end_at": data_end.astimezone(timezone).isoformat()}
    if period_start and period_end:
        return {"start_at": period_start.astimezone(timezone).isoformat(), "end_at": period_end.astimezone(timezone).isoformat()}
    return {}



def _vm_sample_window_label(vms: list[dict[str, Any]], report: dict[str, Any]) -> str:
    starts = [_parse_report_datetime(str(vm.get("window_start_at") or "")) for vm in vms]
    ends = [_parse_report_datetime(str(vm.get("window_end_at") or "")) for vm in vms]
    starts = [value for value in starts if value is not None]
    ends = [value for value in ends if value is not None]
    if starts and ends:
        timezone = _report_timezone(report)
        return f"{min(starts).astimezone(timezone).date().isoformat()} - {max(ends).astimezone(timezone).date().isoformat()}"
    return _period_window_label(report)



def _parse_report_datetime(value: str) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None



def _report_timezone(report: dict[str, Any]) -> ZoneInfo:
    timezone_name = str(report.get("timezone") or "Asia/Shanghai")
    try:
        return ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError:
        return ZoneInfo("Asia/Shanghai")



def _capacity_risk_summary(report: dict[str, Any]) -> str:
    clusters = report.get("clusters") or []
    if not clusters:
        return "暂无集群容量数据，无法判断容量风险。"
    ranked = sorted(clusters, key=_cluster_used_ratio, reverse=True)
    high = [cluster for cluster in ranked if _cluster_used_ratio(cluster) >= 0.8]
    warning = [cluster for cluster in ranked if _cluster_used_ratio(cluster) >= 0.75]
    if high:
        return _risk_summary_sentence(high, "使用率超过 80%，容量风险较高")
    if warning:
        return _risk_summary_sentence(warning, "使用率超过 75%，需要关注容量增长")
    return "当前所有集群暂无明显容量风险。"



def _risk_summary_sentence(clusters: list[dict[str, Any]], suffix: str) -> str:
    names = "、".join(_cluster_name(cluster) for cluster in clusters[:3])
    if len(clusters) > 3:
        names += f" 等 {len(clusters)} 个集群"
    return f"{names} {suffix}。"



def _cluster_name(cluster: dict[str, Any]) -> str:
    labels = cluster.get("labels") or {}
    return str(labels.get("cluster") or labels.get("cluster_id") or "未知集群")



def _cluster_full_name(cluster: dict[str, Any]) -> str:
    labels = cluster.get("labels") or {}
    tower = labels.get("tower") or labels.get("tower_id")
    cluster_name = labels.get("cluster") or labels.get("cluster_id") or "未知集群"
    return f"{tower} / {cluster_name}" if tower else str(cluster_name)



def _tower_scope_label(clusters: list[dict[str, Any]], fallback: str) -> str:
    names = []
    for cluster in clusters:
        labels = cluster.get("labels") or {}
        value = labels.get("tower") or labels.get("tower_id")
        if value and str(value) not in names:
            names.append(str(value))
    if not names:
        return fallback
    if len(names) == 1:
        return names[0]
    return f"全部 Tower（{len(names)} 个）"



def _cluster_scope_label(clusters: list[dict[str, Any]], fallback: str) -> str:
    names = []
    for cluster in clusters:
        name = _cluster_full_name(cluster)
        if name and name not in names:
            names.append(name)
    if not names:
        return fallback
    if len(names) == 1:
        return names[0]
    return f"全部集群（{len(names)} 个）"



def _cluster_used_ratio(cluster: dict[str, Any]) -> float:
    total = _float_or_none(cluster.get("total"))
    current = _float_or_none((cluster.get("forecast") or {}).get("current"))
    if total and total > 0 and current is not None:
        return current / total
    return 0.0



def _risk_level(cluster: dict[str, Any]) -> tuple[str, str]:
    used_ratio = _cluster_used_ratio(cluster)
    forecast = cluster.get("forecast") or {}
    future = _float_or_none(forecast.get("forecast_90d")) or _float_or_none(forecast.get("current")) or 0
    warning = _float_or_none(cluster.get("warning")) or 0
    if used_ratio >= 0.8:
        return "高风险", DANGER
    if warning and future >= warning:
        return "需关注", WARNING
    return "正常", SUCCESS



def _overall_risk_status(clusters: list[dict[str, Any]]) -> tuple[str, str, str]:
    if not clusters:
        return "数据不足", WARNING, "当前导出范围暂无集群容量数据。"
    high = [cluster for cluster in clusters if _cluster_used_ratio(cluster) >= 0.8]
    warning = [cluster for cluster in clusters if _cluster_used_ratio(cluster) >= 0.75 or _exhaustion_days(cluster) <= 180]
    if high:
        return "高风险", DANGER, f"{_cluster_full_name(high[0])} 容量使用率已达到高风险阈值。"
    if warning:
        return "需关注", WARNING, f"{_cluster_full_name(warning[0])} 容量增长或预测耗尽时间需要关注。"
    return "正常", SUCCESS, "当前集群整体运行平稳，暂无明显容量风险。"



def _customer_growth_vms(report: dict[str, Any]) -> list[dict[str, Any]]:
    return _merge_growth_candidates(report.get("window_fastest_growing_vms") or [], report.get("month_fastest_growing_vms") or [])



def _report_vm_count(report: dict[str, Any]) -> int:
    return len(_report_vm_keys(report))



def _report_cluster_vm_counts(report: dict[str, Any]) -> dict[tuple[str, str], int]:
    grouped: dict[tuple[str, str], set[str]] = defaultdict(set)
    for tower_id, cluster_id, vm_id in _report_vm_keys(report):
        grouped[(tower_id, cluster_id)].add(vm_id)
    return {key: len(vm_ids) for key, vm_ids in grouped.items()}



def _cluster_vm_count(cluster: dict[str, Any], fallback: int | None) -> int:
    labels = cluster.get("labels") or {}
    value = cluster.get("vm_count") or cluster.get("virtual_machine_count") or cluster.get("vms")
    try:
        numeric = int(value)
        if numeric >= 0:
            return numeric
    except (TypeError, ValueError):
        pass
    return int(fallback or 0)



def _report_vm_keys(report: dict[str, Any]) -> set[tuple[str, str, str]]:
    keys: set[tuple[str, str, str]] = set()
    for key in ("day_fastest_growing_vms", "window_fastest_growing_vms", "month_fastest_growing_vms", "day_new_vms", "month_new_vms"):
        for vm in report.get(key) or []:
            labels = vm.get("labels") or {}
            vm_key = (
                str(labels.get("tower_id") or ""),
                str(labels.get("cluster_id") or ""),
                str(labels.get("vm_id") or labels.get("vm") or labels.get("vm_name") or ""),
            )
            if any(vm_key):
                keys.add(vm_key)
    return keys



def _customer_key_findings(clusters: list[dict[str, Any]], top_vms: list[dict[str, Any]], risk_note: str, context: dict[str, Any]) -> list[str]:
    if not clusters:
        return ["当前导出范围暂无可用于分析的集群容量数据，建议确认 Tower 与集群采集状态。"]
    total_current = sum(float((cluster.get("forecast") or {}).get("current") or 0) for cluster in clusters)
    total_capacity = sum(float(cluster.get("total") or 0) for cluster in clusters)
    total_window = sum(_cluster_period_growth(cluster, _customer_window_days(context["report"])) for cluster in clusters)
    total_forecast_90 = sum(float((cluster.get("forecast") or {}).get("forecast_90d") or 0) for cluster in clusters)
    profile = context["profile"]
    findings = [
        f"当前导出范围已用容量 {_bytes_label(total_current)}，{profile.window_growth_label} {_bytes_label(total_window)}，{risk_note}",
    ]
    if total_capacity > 0:
        findings.append(f"按当前增长趋势推算，90 天后容量预计为 {_bytes_label(total_forecast_90)}，约占总容量 {_bytes_label(total_capacity)} 的 {_percent_label(total_forecast_90 / total_capacity)}。")
    if top_vms:
        first = top_vms[0]
        findings.append(
            f"{profile.display_label}内，{_vm_scope_name(first)} 下的 {_vm_display_name(first)} "
            f"增长最为显著，增长量 {_bytes_label(first.get('growth_amount'))}，建议确认增长来源的合理性。"
        )
    largest_vm = _largest_vm(top_vms)
    if largest_vm:
        findings.append(
            f"{_vm_scope_name(largest_vm)} 下的 {_vm_display_name(largest_vm)} "
            f"当前容量 {_bytes_label((largest_vm.get('forecast') or {}).get('current'))}，建议纳入重点监控清单。"
        )
    high_ratio = [vm for vm in top_vms if float(vm.get("growth_ratio") or 0) >= 1.0]
    if high_ratio:
        findings.append(f"增长率超过 100% 的 VM 有 {len(high_ratio)} 台，可能属于基数较小导致的增长率偏高，需结合绝对值综合评估。")
    return findings[:5]



def _customer_risk_matrix_rows(clusters: list[dict[str, Any]], top_vms: list[dict[str, Any]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    if clusters:
        worst = max(clusters, key=_cluster_used_ratio)
        ratio = _cluster_used_ratio(worst)
        level = "高" if ratio >= 0.8 else "中" if ratio >= 0.75 else "低"
        rows.append({
            "item": "集群整体容量",
            "status": _risk_level(worst)[0],
            "level": level,
            "description": f"{_cluster_full_name(worst)} 当前使用率 {_percent_label(ratio)}，90 天预测 {_bytes_label((worst.get('forecast') or {}).get('forecast_90d'))}。",
        })
        exhaustion_candidates = [
            (_exhaustion_days(cluster), cluster)
            for cluster in clusters
            if _exhaustion_days(cluster) < float("inf")
        ]
        soonest = min(exhaustion_candidates, key=lambda item: item[0]) if exhaustion_candidates else None
        if soonest:
            days, cluster = soonest
            level = "高" if days <= 90 else "中" if days <= 180 else "低"
            rows.append({
                "item": "预计存储耗尽",
                "status": _days_label(days),
                "level": level,
                "description": f"{_cluster_full_name(cluster)} 按当前趋势预计 {_days_label(days)} 后触达容量上限。",
            })
    if top_vms:
        largest = _largest_vm(top_vms)
        if largest:
            rows.append({
                "item": "重点 VM 容量",
                "status": "需关注" if _vm_current_bytes(largest) >= VM_ALERT_BYTES else "正常",
                "level": "中" if _vm_current_bytes(largest) >= VM_ALERT_BYTES else "低",
                "description": f"{_vm_full_name(largest)} 当前容量 {_bytes_label(_vm_current_bytes(largest))}，建议确认是否存在可清理数据。",
            })
        fastest = top_vms[0]
        rows.append({
            "item": "VM 增长来源",
            "status": "需关注" if float(fastest.get("growth_amount") or 0) > 0 else "健康",
            "level": "中" if float(fastest.get("growth_amount") or 0) > 0 else "低",
            "description": f"{_vm_full_name(fastest)} 增长 {_bytes_label(fastest.get('growth_amount'))}，建议确认业务数据增长合理性。",
        })
    if not rows:
        rows.append({"item": "容量数据", "status": "数据不足", "level": "中", "description": "当前缺少可分析的容量或 VM 增长数据。"})
    return rows



def _customer_operation_advice(clusters: list[dict[str, Any]], top_vms: list[dict[str, Any]], profile: ReportPeriodProfile | None = None) -> list[tuple[str, list[str]]]:
    top_vm = top_vms[0] if top_vms else None
    largest = _largest_vm(top_vms)
    high_clusters = [cluster for cluster in clusters if _cluster_used_ratio(cluster) >= 0.8]
    short_items = [
        (
            f"建议确认 {_vm_scope_name(top_vm)} 下的 {_vm_display_name(top_vm)} "
            f"在统计窗口内增长 {_bytes_label(top_vm.get('growth_amount'))} 的业务来源，判断是否为预期写入、日志膨胀或临时数据堆积。"
        )
        if top_vm else "确认当前采集状态和 Prometheus 历史指标完整性，避免因数据缺口影响容量判断。",
        (
            f"对 {_vm_scope_name(largest)} 下的 {_vm_display_name(largest)}（{_bytes_label(_vm_current_bytes(largest))}）"
            "进行存储空间审计，识别可清理的历史数据、快照或日志。"
        )
        if largest else "检查大容量 VM、快照和备份策略，识别可清理空间。",
    ]
    middle_items = [
        "建立月度容量复盘机制，持续跟踪 Top 10 增长最快的虚拟机。",
        "定期复核大容量 VM 和持续增长 VM 的业务归属、数据保留策略和清理窗口，避免单业务长期占用过多集群空间。",
    ]
    long_items = [
        "当容量使用率达到 60% 时启动扩容评估预案，达到 75% 后进入扩容计划跟踪，达到 80% 后优先执行扩容或清理。",
        "定期审查 VM 快照、备份和日志保留策略，避免冗余数据占用有效存储空间。",
    ]
    if high_clusters:
        short_items.insert(0, f"集群 {_cluster_full_name(high_clusters[0])} 已超过 80% 高风险阈值，建议立即确认扩容周期和可清理空间。")
    short_title = (profile or REPORT_PERIOD_PROFILES[30]).short_advice_title
    return [(short_title, short_items), ("中期（1-3 个月）", middle_items), ("三个月以上", long_items)]



def _vm_display_name(vm: dict[str, Any] | None) -> str:
    if not vm:
        return "未知 VM"
    labels = vm.get("labels") or {}
    return str(labels.get("vm") or labels.get("vm_name") or labels.get("vm_id") or "未知 VM")



def _vm_full_name(vm: dict[str, Any] | None) -> str:
    if not vm:
        return "未知 VM"
    labels = vm.get("labels") or {}
    tower = labels.get("tower") or labels.get("tower_id")
    cluster = labels.get("cluster") or labels.get("cluster_id")
    vm_name = _vm_display_name(vm)
    prefix = " / ".join(str(value) for value in [tower, cluster] if value)
    return f"{prefix} / {vm_name}" if prefix else vm_name



def _vm_scope_name(vm: dict[str, Any] | None) -> str:
    if not vm:
        return "未知范围"
    labels = vm.get("labels") or {}
    tower = labels.get("tower") or labels.get("tower_id")
    cluster = labels.get("cluster") or labels.get("cluster_id")
    scope = " / ".join(str(value) for value in [tower, cluster] if value)
    return scope or "未知范围"



def _vm_current_bytes(vm: dict[str, Any] | None) -> float:
    if not vm:
        return 0.0
    return float((vm.get("forecast") or {}).get("current") or 0.0)



def _largest_vm(vms: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not vms:
        return None
    return max(vms, key=_vm_current_bytes)



def _exhaustion_days(cluster: dict[str, Any]) -> float:
    value = (cluster.get("forecast") or {}).get("exhaustion_days")
    try:
        return float(value)
    except (TypeError, ValueError):
        return float("inf")



def _cluster_period_growth(cluster: dict[str, Any], days: int) -> float:
    points = _cluster_points(cluster)
    forecast = cluster.get("forecast") or {}
    current = _float_or_none(forecast.get("current"))
    if current is None and points:
        current = points[-1][1]
    current = current or 0.0
    if len(points) >= 2:
        target = points[-1][0] - days * 86_400
        candidates = [point for point in points if point[0] <= target]
        baseline = candidates[-1][1] if candidates else points[0][1]
        return max(0.0, current - baseline)
    return max(0.0, float(forecast.get("slope_per_day") or 0) * days)



def _cluster_points(cluster: dict[str, Any]) -> list[tuple[int, float]]:
    points = []
    for point in cluster.get("points") or []:
        try:
            points.append((int(float(point[0])), float(point[1])))
        except (TypeError, ValueError, IndexError):
            continue
    return sorted(points)



def _merged_cluster_points(clusters: list[dict[str, Any]]) -> list[tuple[int, float]]:
    by_ts: dict[int, float] = defaultdict(float)
    for cluster in clusters:
        for ts, value in _cluster_points(cluster):
            by_ts[ts] += value
    return sorted(by_ts.items())



def _vms_by_cluster(vms: list[dict[str, Any]]) -> dict[tuple[str, str], list[dict[str, Any]]]:
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for vm in vms:
        grouped[_cluster_key(vm.get("labels", {}))].append(vm)
    return grouped



def _cluster_key(labels: dict[str, Any]) -> tuple[str, str]:
    return (str(labels.get("tower_id") or ""), str(labels.get("cluster_id") or ""))



def _top_vms(vms: list[dict[str, Any]], mode: str) -> list[dict[str, Any]]:
    key = (lambda vm: float(vm.get("growth_ratio") or 0)) if mode == "ratio" else (lambda vm: float(vm.get("growth_amount") or 0))
    return sorted(vms, key=key, reverse=True)



def _merge_growth_candidates(primary: list[dict[str, Any]], secondary: list[dict[str, Any]]) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for vm in [*primary, *secondary]:
        labels = vm.get("labels") or {}
        key = (
            str(labels.get("tower_id") or ""),
            str(labels.get("cluster_id") or ""),
            str(labels.get("vm_id") or labels.get("vm") or ""),
        )
        if key in seen:
            continue
        seen.add(key)
        merged.append(vm)
    return merged



def _float_or_none(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None



def _first_cluster_name(clusters: list[dict[str, Any]]) -> str | None:
    if not clusters:
        return None
    labels = clusters[0].get("labels", {})
    return labels.get("cluster") or labels.get("cluster_id")



def _is_alert_vm(vm: dict[str, Any]) -> bool:
    return float(vm.get("growth_ratio") or 0) > VM_ALERT_RATIO and float(vm.get("growth_amount") or 0) > VM_ALERT_BYTES



def _scope_trend_line_chart(clusters: list[dict[str, Any]], title: str) -> BytesIO | None:
    return _line_chart_image(_merged_cluster_points(clusters), title=title, ylabel="容量 (TiB)")



def _cluster_top_growth_bar_chart(clusters: list[dict[str, Any]], title: str) -> BytesIO | None:
    items = [(_cluster_name(cluster), _cluster_period_growth(cluster, 30)) for cluster in clusters if _cluster_period_growth(cluster, 30) > 0]
    return _horizontal_bar_chart_image(items, title=title, unit="TiB", limit=5, scale="tib")



def _line_chart_image(points: list[tuple[int, float]], title: str, ylabel: str) -> BytesIO | None:
    if len(points) < 2:
        return None
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.dates as mdates
        import matplotlib.pyplot as plt
        _configure_chart_fonts(matplotlib)
        matplotlib.rcParams["font.family"] = [CHART_FONT_FAMILY, "DejaVu Serif"]
        matplotlib.rcParams["axes.unicode_minus"] = False
    except Exception:
        return None
    dates = [datetime.fromtimestamp(ts) for ts, _ in points]
    values = [_bytes_to_tib(value) for _, value in points]
    fig, ax = plt.subplots(figsize=(6.6, 4.0), dpi=180)
    ax.plot(dates, values, color="#003BFF", linewidth=1.4)
    ax.set_title(title, fontsize=12, pad=8)
    ax.set_xlabel("时间", fontsize=10)
    ax.set_ylabel(ylabel, fontsize=10)
    y_min, y_max = _chart_y_limits(values)
    ax.set_ylim(y_min, y_max)
    ax.grid(axis="y", color="#E6EDF5", linewidth=0.6)
    ax.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=4, maxticks=6))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m-%d"))
    ax.tick_params(axis="both", labelsize=9, colors="#333333")
    fig.tight_layout()
    return _figure_bytes(fig, plt)



def _chart_color(value: str) -> str:
    color = value.strip()
    return color if color.startswith("#") else f"#{color}"



def _horizontal_bar_chart_image(items: list[tuple[str, float]], title: str, unit: str, limit: int, scale: str) -> BytesIO | None:
    if not items:
        return None
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        _configure_chart_fonts(matplotlib)
        matplotlib.rcParams["font.family"] = [CHART_FONT_FAMILY, "DejaVu Serif"]
        matplotlib.rcParams["axes.unicode_minus"] = False
    except Exception:
        return None
    ranked = sorted(items, key=lambda item: item[1], reverse=True)[:limit]
    names = [_truncate_label(name) for name, _ in ranked][::-1]
    values = [_chart_bar_value(value, scale) for _, value in ranked][::-1]
    max_value = max(values) if values else 0
    fig, ax = plt.subplots(figsize=(6.8, max(2.2, 0.45 * len(values) + 1.1)), dpi=180)
    bars = ax.barh(names, values, color="#1155CC", height=0.32)
    ax.set_title(title, fontsize=12, pad=8)
    ax.set_xlim(0, max(max_value * 1.18, 1))
    ax.grid(axis="x", color="#E6EDF5", linewidth=0.6)
    for bar, value in zip(bars, values):
        ax.text(bar.get_width(), bar.get_y() + bar.get_height() / 2, f"{value:.2f} {unit}", va="center", ha="left", fontsize=9)
    fig.tight_layout()
    return _figure_bytes(fig, plt)



def _configure_chart_fonts(matplotlib: Any) -> None:
    font_path = Path(CHART_FONT_PATH)
    if not font_path.exists():
        return
    try:
        from matplotlib import font_manager
        font_manager.fontManager.addfont(str(font_path))
    except Exception:
        return



def _figure_bytes(fig: Any, plt: Any) -> BytesIO:
    output = BytesIO()
    fig.savefig(output, format="png", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    output.seek(0)
    return output



def _chart_y_limits(values: list[float]) -> tuple[float, float]:
    if not values:
        return 0.0, 1.0
    y_min = min(values)
    y_max = max(values)
    padding = max((y_max - y_min) * 0.18, y_max * 0.01, 0.1) if y_min != y_max else max(abs(y_max) * 0.05, 0.1)
    lower = max(0.0, y_min - padding)
    upper = y_max + padding
    return lower, upper if upper > lower else lower + 1.0



def _truncate_label(value: str, limit: int = 22) -> str:
    text = str(value)
    return text if len(text) <= limit else f"{text[:8]}...{text[-8:]}"



def _bytes_to_tib(value: Any) -> float:
    try:
        return float(value or 0) / 1024**4
    except (TypeError, ValueError):
        return 0.0



def _bytes_to_gb(value: Any) -> float:
    try:
        return float(value or 0) / 1024**3
    except (TypeError, ValueError):
        return 0.0



def _chart_bar_value(value: float, scale: str) -> float:
    if scale == "gb":
        return _bytes_to_gb(value)
    return _bytes_to_tib(value)



def _local_now(settings: V2Settings) -> datetime:
    try:
        return datetime.now(ZoneInfo(settings.timezone))
    except ZoneInfoNotFoundError:
        return datetime.now()



def _slug(value: str) -> str:
    normalized = "".join(char if char.isalnum() or char in {"-", "_"} else "-" for char in value.strip())
    return normalized.strip("-") or "all"



def _bytes_label(value: Any) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return "-"
    units = ["B", "KB", "MB", "GB", "TB", "PB"]
    index = 0
    while abs(numeric) >= 1024 and index < len(units) - 1:
        numeric /= 1024
        index += 1
    return f"{numeric:.2f} {units[index]}"



def _signed_bytes_label(value: Any) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return "-"
    label = _bytes_label(abs(numeric))
    if numeric > 0:
        return f"+{label}"
    if numeric < 0:
        return f"-{label}"
    return label



def _days_label(value: Any) -> str:
    try:
        return f"{float(value):.0f} 天"
    except (TypeError, ValueError):
        return "未触发"



def _percent_label(value: Any) -> str:
    try:
        return f"{float(value) * 100:.2f}%"
    except (TypeError, ValueError):
        return "-"


# v2 keeps its report data service, task flow and Excel workbook, but the Word
# document intentionally uses a compact customer-facing delivery template.

def _growth_rate_method_lines(report: dict[str, Any]) -> list[str]:
    growth_rate = report.get("cluster_growth_rate") or {}
    return [
        _growth_rate_method_line("日", "最近一天净变化", growth_rate.get("per_day"), growth_rate.get("day_sample_sufficient"), "/天"),
        _growth_rate_method_line("月", "最近 30 天趋势折算", growth_rate.get("per_month"), growth_rate.get("month_sample_sufficient"), "/月"),
        _growth_rate_method_line("季度", "最近 90 天趋势折算", growth_rate.get("per_quarter"), growth_rate.get("quarter_sample_sufficient"), "/季度"),
    ]



def _growth_rate_method_line(label: str, method: str, value: Any, sample_sufficient: Any, unit: str) -> str:
    value_label = _growth_rate_value_label(value, sample_sufficient, unit)
    if value is None:
        return f"{label}：{value_label}；{method}"
    return f"{label}：{method}，{value_label}"



def _growth_rate_value_label(value: Any, sample_sufficient: Any, unit: str) -> str:
    if value is None:
        label = "数据不足"
    else:
        label = f"{_signed_bytes_label(value)}{unit}"
    if sample_sufficient is False:
        label = f"{label}（样本不足）"
    return label



def _customer_window_days(report: dict[str, Any]) -> int:
    try:
        value = int(report.get("window_days") or (report.get("period_window") or {}).get("days") or 30)
    except (TypeError, ValueError):
        value = 30
    return max(1, value)



def _report_data_quality(report: dict[str, Any]) -> dict[str, Any]:
    quality = report.get("data_quality")
    return quality if isinstance(quality, dict) else {"status": "unknown", "messages": ["当前报表未包含数据质量检查结果。"]}



def _data_quality_status_message(quality: dict[str, Any]) -> str:
    status = str(quality.get("status") or "unknown")
    if status == "ok":
        return "当前报表范围内未发现明显数据缺口。"
    if status == "warning":
        return "当前报表存在样本不足、缺采或部分集群样本不完整，趋势与预测结论需结合实际采集窗口理解。"
    if status == "critical":
        return "当前数据质量异常，趋势和预测结果仅供排障参考，不建议直接作为容量决策依据。"
    messages = quality.get("messages") or []
    return str(messages[0]) if messages else "当前报表未包含数据质量检查结果。"



def _data_quality_status_label(quality: dict[str, Any]) -> str:
    status = str(quality.get("status") or "unknown")
    return {"ok": "数据质量正常", "warning": "数据质量需关注", "critical": "数据质量异常"}.get(status, "数据质量未知")



def _data_quality_summary_rows(report: dict[str, Any], quality: dict[str, Any]) -> list[tuple[str, str]]:
    missing_dates = quality.get("missing_collection_dates") or []
    incomplete_clusters = quality.get("incomplete_clusters") or []
    sqlite_vm_count = _int_label(quality.get("sqlite_vm_count"))
    prometheus_vm_count = _int_label(quality.get("prometheus_vm_series_count"))
    sqlite_cluster_count = _int_label(quality.get("sqlite_cluster_count"))
    prometheus_cluster_count = _int_label(quality.get("prometheus_cluster_series_count"))
    diff = quality.get("vm_count_difference")
    ratio = quality.get("vm_count_difference_ratio")
    rows = [
        ("总体状态", _data_quality_status_label(quality)),
        ("本报表统计窗口", _requested_report_window_label(report)),
        ("实际采集窗口", _data_quality_window_label(quality.get("actual_data_window"), report)),
        ("样本是否足够", "是" if quality.get("sample_sufficient") else "否"),
        ("缺采天数", f"{len(missing_dates)} 天" + (f"（{', '.join(map(str, missing_dates[:10]))}）" if missing_dates else "")),
        ("数据不完整集群", f"{len(incomplete_clusters)} 个"),
        ("最近成功采集时间", _datetime_label(quality.get("latest_success_at"), report)),
        ("最新 Prometheus 样本时间", _datetime_label(quality.get("latest_prometheus_sample_at"), report)),
        ("SQLite 当前 VM 数", sqlite_vm_count),
        ("Prometheus 当前 VM series 数", prometheus_vm_count),
        ("差异数量", _int_label(diff)),
        ("差异比例", _percent_label(ratio)),
        ("SQLite 启用集群数", sqlite_cluster_count),
        ("Prometheus 当前集群 series 数", prometheus_cluster_count),
    ]
    return rows



def _data_quality_window_label(window: Any, report: dict[str, Any]) -> str:
    if not isinstance(window, dict):
        return "-"
    start = _datetime_label(window.get("start_at"), report, date_only=True)
    end = _datetime_label(window.get("end_at"), report, date_only=True)
    days = window.get("days")
    if start != "-" and end != "-":
        suffix = f"，约 {days} 天" if days not in (None, "") else ""
        return f"{start} - {end}{suffix}"
    return "-"



def _datetime_label(value: Any, report: dict[str, Any], *, date_only: bool = False) -> str:
    if value in (None, ""):
        return "-"
    parsed = _parse_report_datetime(str(value))
    if not parsed:
        return str(value)
    converted = parsed.astimezone(_report_timezone(report))
    return converted.date().isoformat() if date_only else converted.strftime("%Y-%m-%d %H:%M")



def _int_label(value: Any) -> str:
    try:
        return str(int(value))
    except (TypeError, ValueError):
        return "-"
