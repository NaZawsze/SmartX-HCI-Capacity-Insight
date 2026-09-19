from __future__ import annotations

"""Excel (xlsx) report builder and template writers."""

from copy import copy
from io import BytesIO
from pathlib import Path
from typing import Any
from openpyxl.chart import BarChart, Reference
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from app.v2.config import V2Settings

from .common import (
    ReportPeriodProfile, _capacity_risk_summary, _cluster_full_name, _cluster_key, _cluster_name, _cluster_period_growth, _cluster_scope_label, _cluster_used_ratio,
    _customer_growth_vms, _customer_key_findings, _data_quality_status_message, _data_quality_summary_rows, _days_label, _effective_report_window, _export_context, _float_or_none,
    _growth_rate_method_lines, _is_alert_vm, _line_chart_image, _merged_cluster_points, _overall_risk_status, _parse_report_datetime, _percent_label, _period_window_label, _persist_report, _report_data_quality,
    _risk_level, _top_vms, _tower_scope_label, _vm_sample_window_label, _vms_by_cluster, ACCENT, ACCENT_DARK, ACCENT_LIGHT,
    ACCENT_SOFT, GROWTH_BLUE, RATIO_RED, REPORT_COVER_SUBTITLE, REPORT_COVER_TITLE, REPORT_PRODUCT_NAME, TEXT_DARK, TEXT_MUTED,
    VM_ALERT_FILL, XLSX_CLUSTER_TEMPLATE_SHEET, XLSX_FONT_NAME, XLSX_TEMPLATE_PATH,
)

def build_report_xlsx(report: dict[str, Any], settings: V2Settings, *, period_days: int) -> tuple[bytes, str, Path, str]:
    context = _export_context(report, settings, period_days, "xlsx")
    clusters = report.get("clusters") or []
    month_vms = report.get("month_fastest_growing_vms") or []
    profile = context["profile"]
    workbook = _load_customer_xlsx_template()
    _write_xlsx_template_cover(workbook["封面"], report, context, settings)
    _write_xlsx_template_summary(workbook["执行摘要"], report, context, clusters, settings, profile)
    quality_sheet = _get_or_create_sheet_after(workbook, "数据质量说明", after="执行摘要")
    _write_xlsx_data_quality_sheet(quality_sheet, report, context)
    _write_xlsx_template_capacity_trend(workbook["容量趋势"], report, context, clusters, profile)
    top_sheet = workbook["VM增长TOP20"] if "VM增长TOP20" in workbook.sheetnames else _get_or_create_sheet(workbook, "VM增长TOP100")
    top_sheet.title = "VM增长TOP100"
    _write_xlsx_template_vm_top100(top_sheet, _customer_growth_vms(report), report, profile)
    day_growth_sheet = workbook["日增长详情"]
    day_growth_sheet.sheet_properties.tabColor = None
    _write_xlsx_template_growth_detail(
        day_growth_sheet,
        report.get("day_fastest_growing_vms") or [],
        report,
        title="日增长最快虚拟机",
        growth_header="日增长量",
    )
    month_growth_sheet = _get_or_create_sheet_after(workbook, "月增长详情", after="日增长详情")
    _write_xlsx_template_growth_detail(
        month_growth_sheet,
        report.get("month_fastest_growing_vms") or [],
        report,
        title="月增长最快虚拟机",
        growth_header="月增长量",
    )

    _write_simple_vm_sheet(_get_or_create_sheet(workbook, "本日新建VM"), report.get("day_new_vms") or [], "暂无本日新建 VM", include_growth=False)
    _write_simple_vm_sheet(_get_or_create_sheet(workbook, "本月新建VM"), report.get("month_new_vms") or [], "暂无本月新建 VM", include_growth=False)
    vms_by_cluster = _vms_by_cluster(month_vms)
    for cluster in clusters:
        labels = cluster.get("labels", {})
        sheet = _clone_cluster_template_sheet(
            workbook,
            _safe_sheet_name(labels.get("cluster") or labels.get("cluster_id") or "集群"),
        )
        _write_cluster_vm_sheet(sheet, cluster, vms_by_cluster.get(_cluster_key(labels), []), report, profile)
    _remove_sheets(workbook, ["目录", "范围明细", "集群汇总", "VM_TOP100_汇总", "汇总", XLSX_CLUSTER_TEMPLATE_SHEET])
    _normalize_xlsx_fonts(workbook)

    output = BytesIO()
    workbook.save(output)
    return _persist_report(output.getvalue(), settings, context["filename"])


def _load_customer_xlsx_template() -> Workbook:
    if XLSX_TEMPLATE_PATH.exists():
        return load_workbook(XLSX_TEMPLATE_PATH)
    workbook = Workbook()
    workbook.active.title = "封面"
    for name in ["执行摘要", "容量趋势", "VM增长TOP20", "日增长详情"]:
        workbook.create_sheet(name)
    return workbook


def _get_or_create_sheet(workbook: Workbook, name: str):
    return workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)


def _get_or_create_sheet_after(workbook: Workbook, name: str, *, after: str):
    if name in workbook.sheetnames:
        sheet = workbook[name]
        current_index = workbook._sheets.index(sheet)
        target_index = workbook.sheetnames.index(after) + 1
        if current_index != target_index:
            workbook._sheets.pop(current_index)
            workbook._sheets.insert(target_index, sheet)
        return sheet
    return workbook.create_sheet(name, workbook.sheetnames.index(after) + 1)


def _clone_cluster_template_sheet(workbook: Workbook, name: str):
    if name in workbook.sheetnames:
        workbook.remove(workbook[name])
    if XLSX_CLUSTER_TEMPLATE_SHEET in workbook.sheetnames:
        sheet = workbook.copy_worksheet(workbook[XLSX_CLUSTER_TEMPLATE_SHEET])
        sheet.title = name
        sheet.sheet_state = "visible"
        return sheet
    return workbook.create_sheet(name)


def _remove_sheets(workbook: Workbook, names: list[str]) -> None:
    for name in names:
        if name in workbook.sheetnames and len(workbook.sheetnames) > 1:
            workbook.remove(workbook[name])


def _clear_xlsx_sheet(sheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def _clear_xlsx_tables(sheet) -> None:
    for name in list(sheet.tables.keys()):
        del sheet.tables[name]


def _reset_xlsx_sheet_rows(sheet) -> None:
    for merged_range in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged_range))
    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)


def _set_xlsx_cell(sheet, coordinate: str, value: Any, *, size: float | None = None, bold: bool | None = None, color: str | None = None, align: str | None = None) -> None:
    cell = sheet[coordinate]
    cell.value = value
    cell.font = Font(
        name="Noto Sans CJK SC",
        size=size if size is not None else cell.font.sz,
        bold=bold if bold is not None else cell.font.bold,
        color=color if color is not None else cell.font.color,
    )
    if align is not None:
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)


def _write_xlsx_template_cover(sheet, report: dict[str, Any], context: dict[str, Any], settings: V2Settings) -> None:
    clusters = report.get("clusters") or []
    _clear_xlsx_sheet(sheet)
    _set_xlsx_cell(sheet, "B7", REPORT_PRODUCT_NAME, size=14, color=GROWTH_BLUE, align="center")
    _set_xlsx_cell(sheet, "B8", REPORT_COVER_TITLE, size=26, bold=True, color=ACCENT_DARK, align="center")
    _set_xlsx_cell(sheet, "B9", REPORT_COVER_SUBTITLE, size=12, color=TEXT_MUTED, align="center")
    _set_xlsx_cell(sheet, "B11", "───────────────────────────────────────────────────────", size=7, color="CCCCCC", align="center")
    _set_xlsx_cell(sheet, "B13", "客户名称：", size=12, color="555555", align="center")
    _set_xlsx_cell(sheet, "B14", f"Tower范围：{_tower_scope_label(clusters, context['scope_label'])}", size=12, color="555555", align="center")
    _set_xlsx_cell(sheet, "B15", f"集群范围：{_cluster_scope_label(clusters, context['scope_label'])}", size=12, color="555555", align="center")
    _set_xlsx_cell(sheet, "B16", f"统计窗口：{_period_window_label(report)}", size=12, color="555555", align="center")


def _write_xlsx_template_summary(
    sheet,
    report: dict[str, Any],
    context: dict[str, Any],
    clusters: list[dict[str, Any]],
    settings: V2Settings,
    profile: ReportPeriodProfile,
) -> None:
    _clear_xlsx_sheet(sheet)
    current = sum(float((cluster.get("forecast") or {}).get("current") or 0) for cluster in clusters)
    forecast_90d = sum(float((cluster.get("forecast") or {}).get("forecast_90d") or 0) for cluster in clusters)
    total = sum(float(cluster.get("total") or 0) for cluster in clusters)
    growth = sum(_cluster_period_growth(cluster, profile.days) for cluster in clusters)
    risk_status, _, risk_note = _overall_risk_status(clusters)
    growth_ratio = growth / max(current - growth, 1) if growth > 0 else 0.0
    usage_ratio = current / total if total else 0.0
    forecast_ratio = forecast_90d / total if total else 0.0

    _set_xlsx_cell(sheet, "A1", "一、执行摘要", size=16, bold=True, color=ACCENT_DARK)
    _set_xlsx_cell(
        sheet,
        "A2",
        f"报告周期：{_period_window_label(report)}  |  Tower：{_tower_scope_label(clusters, context['scope_label'])}  |  集群：{_cluster_scope_label(clusters, context['scope_label'])}  |  统计口径：{profile.window_growth_label}",
        size=10,
        color=TEXT_MUTED,
    )
    headers = ["当前已用容量", profile.window_growth_label, "90 天预测容量", "风险状态"]
    values = [_xlsx_bytes_label(current), _xlsx_signed_bytes_label(growth), _xlsx_bytes_label(forecast_90d), risk_status]
    subtitles = [f"使用率 {_percent_label(usage_ratio)}", f"增长率 {_percent_label(growth_ratio)}", f"预计使用率 {_percent_label(forecast_ratio)}", risk_note]
    for column, (header, value, subtitle) in enumerate(zip(headers, values, subtitles), start=1):
        for row in [4, 5, 6]:
            cell = sheet.cell(row=row, column=column)
            cell.fill = PatternFill(fill_type="solid", fgColor=ACCENT_LIGHT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.cell(row=4, column=column).value = header
        sheet.cell(row=4, column=column).font = Font(name=XLSX_FONT_NAME, size=10, bold=True, color="000000")
        sheet.cell(row=5, column=column).value = value
        sheet.cell(row=5, column=column).font = Font(name=XLSX_FONT_NAME, size=18, bold=True, color=RATIO_RED if risk_status == "高风险" and column == 4 else ACCENT_DARK)
        sheet.cell(row=6, column=column).value = subtitle
        sheet.cell(row=6, column=column).font = Font(name=XLSX_FONT_NAME, size=10, bold=True, color="000000")

    _set_xlsx_cell(sheet, "A8", "关键发现", size=14, bold=True, color=ACCENT_DARK)
    findings = _customer_key_findings(clusters, _customer_growth_vms(report), risk_note, context)
    for offset, finding in enumerate(findings[:4], start=9):
        _set_xlsx_cell(sheet, f"A{offset}", finding, size=11, color=TEXT_DARK)
    _set_xlsx_cell(sheet, "A13", "容量风险摘要", size=14, bold=True, color=ACCENT_DARK)
    _set_xlsx_cell(sheet, "A14", _capacity_risk_summary(report), size=11, color=TEXT_DARK)
    _set_xlsx_cell(sheet, "A15", f"当前软件版本：{settings.app_version}", size=10, color=TEXT_MUTED)
    _set_xlsx_cell(sheet, "A16", _profile_sample_notice(report, profile), size=10, color=TEXT_MUTED)
    _set_xlsx_cell(sheet, "A18", "容量增长速率口径", size=14, bold=True, color=ACCENT_DARK)
    growth_lines = list(_growth_rate_method_lines(report))
    for offset, line in enumerate(growth_lines, start=19):
        _set_xlsx_cell(sheet, f"A{offset}", line, size=10, color=TEXT_DARK)
    disclaimer_row = 19 + len(growth_lines) + 1
    _set_xlsx_cell(
        sheet,
        f"A{disclaimer_row}",
        "声明：预测结果基于历史增长趋势推算，预测值可能会有偏差，仅供参考，请以实际使用情况为准。",
        size=10,
        color=TEXT_MUTED,
    )
    for row in [1, 2, 8, 9, 10, 11, 12, 13, 14, 15, 16, 18, 19, 20, 21, disclaimer_row]:
        _merge_title_row(sheet, row, 1, 6)
    for column, width in {"A": 50, "B": 20.5, "C": 18, "D": 38.83203125, "E": 16}.items():
        sheet.column_dimensions[column].width = width
    for row, height in {
        1: 30,
        2: 38,
        3: 15,
        4: 24,
        5: 36,
        6: 28,
        7: 15,
        8: 30,
        9: 38,
        10: 38,
        11: 38,
        12: 38,
        13: 15,
        14: 15,
        15: 15,
        16: 38,
        18: 26,
        19: 28,
        20: 28,
        21: 28,
    }.items():
        sheet.row_dimensions[row].height = height


def _write_xlsx_template_capacity_trend(sheet, report: dict[str, Any], context: dict[str, Any], clusters: list[dict[str, Any]], profile: ReportPeriodProfile) -> None:
    _reset_xlsx_sheet_rows(sheet)
    sheet.append(["二、集群容量趋势"])
    sheet.append([f"统计窗口：{_period_window_label(report)}；{_profile_sample_notice(report, profile)}"])
    headers = ["Tower", "集群", "当前已用容量", profile.window_growth_label, "90 天预测容量", "总容量", "使用率", "预计耗尽天数", "风险状态"]
    sheet.append(headers)
    if not clusters:
        sheet.append(["当前范围暂无集群容量数据"] + [""] * (len(headers) - 1))
    for cluster in clusters:
        labels = cluster.get("labels") or {}
        forecast = cluster.get("forecast") or {}
        risk, _ = _risk_level(cluster)
        sheet.append(
            [
                labels.get("tower") or labels.get("tower_id") or "",
                labels.get("cluster") or labels.get("cluster_id") or "",
                _xlsx_bytes_label(forecast.get("current")),
                _xlsx_signed_bytes_label(_cluster_period_growth(cluster, profile.days)),
                _xlsx_bytes_label(forecast.get("forecast_90d")),
                _xlsx_bytes_label(cluster.get("total")),
                _percent_label(_cluster_used_ratio(cluster)),
                _days_label(forecast.get("exhaustion_days")),
                risk,
            ]
        )
    _style_customer_xlsx_table(sheet, title_rows={1, 2}, header_rows={3})
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{sheet.max_row}"
    for column, width in {
        "A": 50,
        "B": 25.83203125,
        "C": 16.83203125,
        "D": 20.83203125,
        "E": 16.83203125,
        "F": 14.83203125,
        "G": 16.83203125,
        "I": 14.83203125,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 51
    sheet.row_dimensions[3].height = 16
    for row_index in range(4, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 30

    _add_capacity_trend_chart(sheet, clusters)
    _setup_xlsx_print_layout(sheet)


def _add_capacity_trend_chart(sheet, clusters: list[dict[str, Any]]) -> None:
    points = _merged_cluster_points(clusters)
    image = _line_chart_image(points, "集群容量使用趋势", "容量 (TiB)")
    if image is None:
        return
    anchor_row = sheet.max_row + 2
    xlsx_image = XlsxImage(image)
    xlsx_image.width = 660
    xlsx_image.height = 400
    sheet.add_image(xlsx_image, f"A{anchor_row}")


def _setup_xlsx_print_layout(sheet) -> None:
    sheet.print_area = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _write_xlsx_data_quality_sheet(sheet, report: dict[str, Any], context: dict[str, Any]) -> None:
    _reset_xlsx_sheet_rows(sheet)
    quality = _report_data_quality(report)
    summary_rows = _data_quality_summary_rows(report, quality)
    sheet.append(["数据质量说明"])
    sheet.append([_data_quality_status_message(quality)])
    sheet.append(["指标", "说明"])
    for label, value in summary_rows:
        sheet.append([label, value])

    incomplete_clusters = quality.get("incomplete_clusters") or []
    start_row = sheet.max_row + 2
    sheet.cell(row=start_row, column=1).value = "数据不完整集群"
    sheet.cell(row=start_row + 1, column=1).value = "Tower"
    sheet.cell(row=start_row + 1, column=2).value = "集群"
    sheet.cell(row=start_row + 1, column=3).value = "原因"
    if incomplete_clusters:
        for item in incomplete_clusters:
            sheet.append([
                item.get("tower") or item.get("tower_id") or "-",
                item.get("cluster") or item.get("cluster_id") or "-",
                item.get("reason") or "-",
            ])
    else:
        sheet.append(["-", "-", "当前报表范围内未发现明显数据不完整集群"])

    _style_customer_xlsx_table(sheet, title_rows={1, 2, start_row}, header_rows={3, start_row + 1})
    _merge_title_row(sheet, 1, 1, 3)
    _merge_title_row(sheet, 2, 1, 3)
    _merge_title_row(sheet, start_row, 1, 3)
    sheet.freeze_panes = "A4"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 44
    sheet.column_dimensions["C"].width = 58
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 52
    for row_index in range(4, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 28


def _write_xlsx_template_vm_top100(sheet, vms: list[dict[str, Any]], report: dict[str, Any], profile: ReportPeriodProfile) -> None:
    _reset_xlsx_sheet_rows(sheet)
    window_label = _vm_sample_window_label(vms, report)
    sheet.append([f"三、{profile.vm_growth_title}全部虚拟机（{window_label}）", "", "", "", "", "", "", "", "四、虚拟机增长率全部虚拟机"])
    sheet.append([])
    sheet.append(["按增长量降序", "", "", "", "", "", "", "", "按增长率降序"])
    amount_headers = ["排名", "虚拟机名称", "当前容量", "期初容量", "增长量", "增长率", "风险"]
    ratio_headers = ["排名", "虚拟机名称", "当前容量", "期初容量", "增长量", "增长率"]
    for index, value in enumerate(amount_headers, start=1):
        sheet.cell(row=4, column=index).value = value
    for index, value in enumerate(ratio_headers, start=9):
        sheet.cell(row=4, column=index).value = value
    amount_vms = _top_vms(vms, "amount")
    ratio_vms = _top_vms(vms, "ratio")
    if not amount_vms and not ratio_vms:
        sheet.cell(row=5, column=1).value = profile.vm_empty_text
    for row_index, vm in enumerate(amount_vms, start=5):
        for column, value in enumerate(_xlsx_vm_display_row(vm, row_index - 4, include_risk=True), start=1):
            sheet.cell(row=row_index, column=column).value = value
    for row_index, vm in enumerate(ratio_vms, start=5):
        for column, value in enumerate(_xlsx_vm_display_row(vm, row_index - 4, include_risk=False), start=9):
            sheet.cell(row=row_index, column=column).value = value
    _style_customer_xlsx_table(sheet, title_rows={1, 3}, header_rows={4})
    _apply_vm_top100_layout(sheet, left_header_row=4, right_header_row=4)
    sheet.freeze_panes = "A5"
    sheet.sheet_view.topLeftCell = "A1"
    for selection in sheet.sheet_view.selection:
        selection.activeCell = "A5"
        selection.sqref = "A5"


def _write_xlsx_template_growth_detail(
    sheet,
    vms: list[dict[str, Any]],
    report: dict[str, Any],
    *,
    title: str,
    growth_header: str,
) -> None:
    _reset_xlsx_sheet_rows(sheet)
    sheet.append([f"{title}（{_period_window_label(report)}）"])
    sheet.append([])
    headers = ["排名", "虚拟机名称", "Tower", "集群", "当前容量", "期初容量", growth_header, "增长率", "风险"]
    sheet.append(headers)
    if not vms:
        sheet.append(["暂无日增长 VM 数据"])
    for index, vm in enumerate(vms, start=1):
        row = _xlsx_vm_display_row(vm, index, include_risk=True)
        labels = vm.get("labels") or {}
        sheet.append([row[0], row[1], labels.get("tower") or labels.get("tower_id") or "", labels.get("cluster") or labels.get("cluster_id") or "", *row[2:]])
    _style_customer_xlsx_table(sheet, title_rows={1}, header_rows={3})
    _merge_title_row(sheet, 1, 1, 9)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{sheet.max_row}"
    for column, width in {
        "A": 15.83203125,
        "B": 44.5,
        "C": 25.83203125,
        "E": 14.83203125,
        "H": 10.83203125,
        "I": 10,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[3].height = 16
    for row_index in range(4, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 30


def _style_customer_xlsx_table(sheet, *, title_rows: set[int] | None = None, header_rows: set[int] | None = None) -> None:
    title_rows = title_rows or set()
    header_rows = header_rows or {1}
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.font = Font(name=XLSX_FONT_NAME, size=11, color=TEXT_DARK)
            if cell.row in title_rows:
                cell.font = Font(name=XLSX_FONT_NAME, size=16 if cell.row == 1 else 12, bold=True, color=ACCENT_DARK)
            if cell.row in header_rows:
                cell.font = Font(name=XLSX_FONT_NAME, size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor=ACCENT_DARK)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif cell.row not in title_rows and cell.row % 2 == 0:
                cell.fill = PatternFill(fill_type="solid", fgColor=ACCENT_SOFT)


def _merge_title_row(sheet, row: int, start_col: int, end_col: int) -> None:
    if end_col <= start_col:
        return
    range_ref = f"{get_column_letter(start_col)}{row}:{get_column_letter(end_col)}{row}"
    if range_ref not in {str(merged) for merged in sheet.merged_cells.ranges}:
        sheet.merge_cells(range_ref)
    cell = sheet.cell(row=row, column=start_col)
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def _apply_vm_top100_layout(sheet, *, left_header_row: int, right_header_row: int | None = None) -> None:
    widths = {
        "A": 15.83203125,
        "B": 50,
        "C": 16,
        "F": 12,
        "H": 10,
        "I": 15.83203125,
        "J": 50,
        "K": 16,
        "N": 12,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row_index in range(1, sheet.max_row + 1):
        if row_index in {1, 3}:
            sheet.row_dimensions[row_index].height = 34
        elif row_index in {left_header_row, right_header_row}:
            sheet.row_dimensions[row_index].height = 28
        else:
            sheet.row_dimensions[row_index].height = 15
    _merge_title_row(sheet, 1, 1, 7)
    if sheet.cell(row=1, column=9).value:
        _merge_title_row(sheet, 1, 9, 14)
    if sheet.cell(row=3, column=1).value:
        _merge_title_row(sheet, 3, 1, 7)
    if sheet.cell(row=3, column=9).value:
        _merge_title_row(sheet, 3, 9, 14)
    if sheet.max_row >= 1:
        sheet.merge_cells(start_row=1, start_column=8, end_row=sheet.max_row, end_column=8)
        separator = sheet.cell(row=1, column=8)
        separator.value = None
        separator.fill = PatternFill(fill_type=None)
        sheet.column_dimensions["H"].width = 10


def _xlsx_vm_display_row(vm: dict[str, Any], rank: int, *, include_risk: bool) -> list[Any]:
    labels = vm.get("labels") or {}
    forecast = vm.get("forecast") or {}
    values: list[Any] = [
        rank,
        labels.get("vm") or labels.get("vm_name") or labels.get("vm_id") or "",
        _xlsx_bytes_label(forecast.get("current")),
        _xlsx_bytes_label(vm.get("previous_value")),
        _xlsx_signed_bytes_label(vm.get("growth_amount")),
        _percent_label(vm.get("growth_ratio")),
    ]
    if include_risk:
        values.append(_vm_risk_label(vm))
    return values


def _vm_risk_label(vm: dict[str, Any]) -> str:
    if _is_alert_vm(vm):
        return "高"
    ratio = _float_or_none(vm.get("growth_ratio")) or 0.0
    amount = _float_or_none(vm.get("growth_amount")) or 0.0
    if ratio >= 0.1 or amount >= 50 * 1024**3:
        return "中"
    return "低"


def _xlsx_bytes_label(value: Any) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return "-"
    units = ["B", "KiB", "MiB", "GiB", "TiB", "PiB"]
    index = 0
    while abs(numeric) >= 1024 and index < len(units) - 1:
        numeric /= 1024
        index += 1
    return f"{numeric:.2f} {units[index]}"


def _xlsx_signed_bytes_label(value: Any) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return "-"
    label = _xlsx_bytes_label(abs(numeric))
    if numeric > 0:
        return f"+{label}"
    if numeric < 0:
        return f"-{label}"
    return label


def _write_simple_vm_sheet(sheet, vms: list[dict[str, Any]], empty_text: str, *, include_growth: bool) -> None:
    _clear_xlsx_sheet(sheet)
    _clear_xlsx_tables(sheet)
    headers = ["Tower", "集群", "VM", "当前容量"]
    if include_growth:
        headers.extend(["期初容量", "增长量", "增长率"])
    else:
        headers.append("首次出现时间")
    sheet.cell(row=1, column=1).value = sheet.title
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=2, column=column).value = header
    if not vms:
        sheet.cell(row=3, column=1).value = empty_text
    for row_index, vm in enumerate(vms, start=3):
        labels = vm.get("labels", {})
        row = [
            labels.get("tower") or labels.get("tower_id") or "",
            labels.get("cluster") or labels.get("cluster_id") or "",
            labels.get("vm") or labels.get("vm_id") or "",
            _xlsx_bytes_label((vm.get("forecast") or {}).get("current")),
        ]
        if include_growth:
            row.extend([
                _xlsx_bytes_label(vm.get("previous_value")),
                _xlsx_signed_bytes_label(vm.get("growth_amount")),
                _percent_label(vm.get("growth_ratio")),
            ])
        else:
            row.append(vm.get("first_seen_at") or "")
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column).value = value
    _style_customer_xlsx_table(sheet, title_rows={1}, header_rows={2})
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{sheet.max_row}"
    widths = {"A": 17.5, "B": 18, "C": 36, "D": 16}
    widths["G" if include_growth else "E"] = 13 if include_growth else 30
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[2].height = 16
    for row_index in range(3, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 15


def _write_cluster_vm_sheet(sheet, cluster: dict[str, Any], vms: list[dict[str, Any]], report: dict[str, Any], profile: ReportPeriodProfile) -> None:
    _clear_xlsx_tables(sheet)
    _reset_xlsx_sheet_rows(sheet)
    forecast = cluster.get("forecast") or {}
    risk, _ = _risk_level(cluster)
    labels = cluster.get("labels") or {}
    sheet.append([_cluster_full_name(cluster)])
    sheet.append(["Tower", "集群", "当前容量", profile.window_growth_label, "风险", "90 天预测容量", "总容量", "使用率", "预计耗尽天数"])
    sheet.append(
        [
            labels.get("tower") or labels.get("tower_id") or "",
            labels.get("cluster") or labels.get("cluster_id") or "",
            _xlsx_bytes_label(forecast.get("current")),
            _xlsx_signed_bytes_label(_cluster_period_growth(cluster, profile.days)),
            risk,
            _xlsx_bytes_label(forecast.get("forecast_90d")),
            _xlsx_bytes_label(cluster.get("total")),
            _percent_label(_cluster_used_ratio(cluster)),
            _days_label(forecast.get("exhaustion_days")),
        ]
    )
    sheet.append([_profile_sample_notice(report, profile)])
    sheet.append([])
    headers = ["VM", "当前容量", "期初容量", "增长量", "增长率"]
    window_label = _vm_sample_window_label(vms, report)
    sheet.append([f"{profile.vm_growth_title} 增长量全部虚拟机（按增长量降序，统计窗口：{window_label}）"])
    sheet.append(headers)
    amount_start = sheet.max_row + 1
    if not vms:
        sheet.append([profile.vm_empty_text])
    else:
        for vm in _top_vms(vms, "amount"):
            sheet.append(_vm_xlsx_row(vm, include_cluster=False))
    amount_end = sheet.max_row
    sheet.append([])
    sheet.append([f"{profile.vm_growth_title} 增长率全部虚拟机（按增长率降序，统计窗口：{window_label}）"])
    ratio_title_row = sheet.max_row
    sheet.append(headers)
    ratio_header_row = sheet.max_row
    ratio_start = sheet.max_row + 1
    if not vms:
        sheet.append([profile.vm_empty_text])
    else:
        for vm in _top_vms(vms, "ratio"):
            sheet.append(_vm_xlsx_row(vm, include_cluster=False))
    ratio_end = sheet.max_row
    _style_customer_xlsx_table(sheet, title_rows={1, 4, 6, ratio_title_row}, header_rows={2, 7, ratio_header_row})
    _style_vm_rows(sheet, amount_start, amount_end, 5)
    _style_vm_rows(sheet, ratio_start, ratio_end, 5)
    _merge_title_row(sheet, 1, 1, 9)
    _merge_title_row(sheet, 4, 1, 9)
    _merge_title_row(sheet, 6, 1, 5)
    _merge_title_row(sheet, ratio_title_row, 1, 5)
    _apply_cluster_sheet_layout(sheet, amount_header_row=7, ratio_header_row=ratio_header_row)
    sheet.freeze_panes = "A7"
    _add_excel_table(sheet, _table_safe_name(sheet.title, "Amount"), 7, amount_end, len(headers))
    _add_excel_table(sheet, _table_safe_name(sheet.title, "Ratio"), ratio_header_row, ratio_end, len(headers))


def _apply_cluster_sheet_layout(sheet, *, amount_header_row: int, ratio_header_row: int) -> None:
    widths = {
        "A": 36.83203125,
        "B": 39.5,
        "C": 22.6640625,
        "D": 19.6640625,
        "E": 15,
        "F": 16,
        "I": 12,
        "J": 8.83203125,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    amount_start = amount_header_row + 1
    amount_end = ratio_header_row - 2
    ratio_start = ratio_header_row + 1
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[3].height = 24
    sheet.row_dimensions[4].height = 32
    sheet.row_dimensions[6].height = 32
    sheet.row_dimensions[amount_header_row].height = 24
    sheet.row_dimensions[ratio_header_row - 1].height = 32
    sheet.row_dimensions[ratio_header_row].height = 24
    for row_index in range(amount_start, amount_end + 1):
        sheet.row_dimensions[row_index].height = 17
    for row_index in range(ratio_start, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 17

    for cell in sheet[1]:
        if not isinstance(cell, MergedCell):
            cell.font = Font(name=XLSX_FONT_NAME, size=23, bold=True, color=ACCENT_DARK)
    for cell in sheet[3]:
        if not isinstance(cell, MergedCell):
            cell.font = Font(name=XLSX_FONT_NAME, size=12, color=TEXT_DARK)
    for cell in sheet[4]:
        if not isinstance(cell, MergedCell):
            cell.font = Font(name=XLSX_FONT_NAME, size=18, color=ACCENT_DARK)
    for row_index in [6, ratio_header_row - 1]:
        for cell in sheet[row_index]:
            if not isinstance(cell, MergedCell):
                cell.font = Font(name=XLSX_FONT_NAME, size=14, color=ACCENT_DARK)
    for row_index in [*range(amount_start, amount_end + 1), *range(ratio_start, sheet.max_row + 1)]:
        for cell in sheet[row_index]:
            if not isinstance(cell, MergedCell):
                cell.font = Font(name=XLSX_FONT_NAME, size=12, color=TEXT_DARK)
                cell.fill = PatternFill(fill_type=None)


def _normalize_xlsx_fonts(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                font = copy(cell.font)
                font.name = XLSX_FONT_NAME
                font.sz = float(font.sz or 11)
                font.scheme = None
                cell.font = font


def _vm_xlsx_row(vm: dict[str, Any], *, include_cluster: bool) -> list[Any]:
    labels = vm.get("labels", {})
    forecast = vm.get("forecast", {})
    row: list[Any] = []
    if include_cluster:
        row.extend([labels.get("tower") or labels.get("tower_id") or "", labels.get("cluster") or labels.get("cluster_id") or ""])
    row.extend([
        labels.get("vm") or labels.get("vm_name") or labels.get("vm_id") or "",
        _xlsx_bytes_label(forecast.get("current")),
        _xlsx_bytes_label(vm.get("previous_value")),
        _xlsx_signed_bytes_label(vm.get("growth_amount")),
        _percent_label(vm.get("growth_ratio")),
    ])
    return row


def _style_vm_rows(sheet, start_row: int, end_row: int, column_count: int) -> None:
    if end_row < start_row:
        return
    fill = PatternFill(fill_type="solid", fgColor=VM_ALERT_FILL)
    for row_index in range(start_row, end_row + 1):
        values = [sheet.cell(row=row_index, column=column).value for column in range(1, column_count + 1)]
        amount = _parse_xlsx_bytes_label(values[-2]) if len(values) >= 2 else 0
        ratio = _parse_percent_label(values[-1]) if values else 0
        if _is_alert_vm({"growth_amount": amount, "growth_ratio": ratio}):
            for column in range(1, column_count + 1):
                sheet.cell(row=row_index, column=column).fill = fill


def _parse_xlsx_bytes_label(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip().replace(",", "")
    sign = -1.0 if text.startswith("-") else 1.0
    text = text.lstrip("+-").strip()
    parts = text.split()
    if not parts:
        return 0.0
    try:
        number = float(parts[0])
    except (TypeError, ValueError):
        return 0.0
    unit = parts[1] if len(parts) > 1 else "B"
    multipliers = {
        "B": 1,
        "KiB": 1024,
        "MiB": 1024**2,
        "GiB": 1024**3,
        "TiB": 1024**4,
        "PiB": 1024**5,
        "KB": 1024,
        "MB": 1024**2,
        "GB": 1024**3,
        "TB": 1024**4,
        "PB": 1024**5,
    }
    return sign * number * multipliers.get(unit, 1)


def _parse_percent_label(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip()
    if not text.endswith("%"):
        return 0.0
    try:
        return float(text[:-1]) / 100
    except (TypeError, ValueError):
        return 0.0


def _add_excel_table(sheet, name: str, header_row: int, end_row: int, column_count: int) -> None:
    if end_row <= header_row:
        return
    ref = f"A{header_row}:{get_column_letter(column_count)}{end_row}"
    table = Table(displayName=name[:255], ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def _safe_sheet_name(value: str) -> str:
    cleaned = "".join(char if char not in "[]:*?/\\'" else "_" for char in str(value))
    return (cleaned[:31] or "Sheet").strip()


def _table_safe_name(sheet_name: str, suffix: str) -> str:
    return "".join(char for char in f"{sheet_name}_{suffix}" if char.isalnum())[:200] or f"Table{suffix}"


def _profile_sample_notice(report: dict[str, Any], profile: ReportPeriodProfile) -> str:
    effective = _effective_report_window(report)
    start = _parse_report_datetime(str(effective.get("start_at") or ""))
    end = _parse_report_datetime(str(effective.get("end_at") or ""))
    if start and end:
        actual_days = max(1, (end.date() - start.date()).days + 1)
        if actual_days < profile.days:
            return f"说明：当前实际采集样本约 {actual_days} 天，{profile.window_growth_label} 已按当前可用样本窗口计算。"
    return f"说明：{profile.window_growth_label} 按当前统计窗口样本计算；如采集历史不足对应天数，则按可用样本窗口计算。"
