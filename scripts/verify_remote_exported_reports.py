import io
import json
import sys
import urllib.request
import zipfile
from openpyxl import load_workbook

import os
BASE_URL = os.environ.get("SMARTX_API_URL") or (sys.argv[1] if len(sys.argv) > 1 else "http://10.20.11.3:8000")

def main():
    print("1. 正在登录平台...")
    login_data = json.dumps({"username": "admin", "password": "password"}).encode("utf-8")
    req = urllib.request.Request(f"{BASE_URL}/api/auth/login", data=login_data, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode("utf-8"))
        token = res["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print("登录成功！")

    print("\n2. 正在导出 Word (docx) 报表 (period_days=30)...")
    req_word = urllib.request.Request(f"{BASE_URL}/api/reports/export/word?period_days=30", headers=headers)
    with urllib.request.urlopen(req_word) as resp:
        word_bytes = resp.read()
    with open("/tmp/sample.docx", "wb") as f:
        f.write(word_bytes)
    print(f"Word 导出成功，大小: {len(word_bytes)} 字节")

    # 验证 Word XML 内容
    with zipfile.ZipFile(io.BytesIO(word_bytes)) as z:
        word_xml = z.read("word/document.xml").decode("utf-8")
    assert "总已分配容量" in word_xml, "Word 范围信息表缺失「总已分配容量」！"
    assert "已分配（占比）" in word_xml, "Word 集群增长表缺失「已分配（占比）」列！"
    assert "分配空间" in word_xml, "Word 虚拟机增长表缺失「分配空间」列！"
    assert "使用率" in word_xml, "Word 虚拟机增长表缺失「使用率」列！"
    print(" [Word 验证通过] 范围基本信息表包含「总已分配容量」，集群增长表包含「已分配（占比）」，VM 表包含「分配空间」与「使用率」！")

    print("\n3. 正在导出 Excel (xlsx) 报表 (period_days=30)...")
    req_excel = urllib.request.Request(f"{BASE_URL}/api/reports/export/excel?period_days=30", headers=headers)
    with urllib.request.urlopen(req_excel) as resp:
        excel_bytes = resp.read()
    with open("/tmp/sample.xlsx", "wb") as f:
        f.write(excel_bytes)
    print(f"Excel 导出成功，大小: {len(excel_bytes)} 字节")

    # 验证 Excel 内容与布局
    wb = load_workbook(io.BytesIO(excel_bytes))

    # 3.1 执行摘要
    summary = wb["执行摘要"]
    assert summary["B4"].value == "已分配容量", f"执行摘要 B4 预期 '已分配容量'，实际 '{summary['B4'].value}'"
    print(f" [Excel 执行摘要] B4={summary['B4'].value}, B5={summary['B5'].value}, B6={summary['B6'].value}")
    assert summary["E4"].value == "风险状态", f"执行摘要 E4 预期 '风险状态'，实际 '{summary['E4'].value}'"
    assert "虚拟机容量口径说明" in str(summary["A17"].value), f"执行摘要 A17 缺失口径说明，实际 '{summary['A17'].value}'"
    print(f" [Excel 执行摘要] A17 口径说明: {summary['A17'].value}")

    # 3.2 容量趋势
    trend = wb["容量趋势"]
    headers_trend = [trend.cell(row=3, column=col).value for col in range(1, 12)]
    print(f" [Excel 容量趋势] 表头 11 列: {headers_trend}")
    assert "总容量" in headers_trend, "容量趋势缺失总容量！"
    assert "已分配容量" in headers_trend, "容量趋势缺失已分配容量！"
    assert "已分配率" in headers_trend, "容量趋势缺失已分配率！"
    col_total = headers_trend.index("总容量") + 1
    col_alloc = headers_trend.index("已分配容量") + 1
    col_alloc_ratio = headers_trend.index("已分配率") + 1
    assert col_alloc == col_total + 1, "已分配容量应紧随总容量之后！"
    assert col_alloc_ratio == col_alloc + 1, "已分配率应紧随已分配容量之后！"

    # 3.3 TOP100 Sheet
    top_sheet = wb["VM增长TOP100"]
    top_headers_left = [top_sheet.cell(row=4, column=col).value for col in range(1, 12)]
    print(f" [Excel VM增长TOP100] 左表头 11 列: {top_headers_left}")
    assert top_headers_left == ["排名", "虚拟机名称", "所有者", "虚拟机组", "分配空间", "当前容量", "使用率", "期初容量", "增长量", "增长率", "风险"], f"TOP100 左表头不匹配: {top_headers_left}"
    top_headers_right = [top_sheet.cell(row=4, column=col).value for col in range(13, 23)]
    print(f" [Excel VM增长TOP100] 右表头 10 列: {top_headers_right}")
    assert top_headers_right == ["排名", "虚拟机名称", "所有者", "虚拟机组", "分配空间", "当前容量", "使用率", "期初容量", "增长量", "增长率"], f"TOP100 右表头不匹配: {top_headers_right}"
    # 打印第一台 VM 的分配空间与使用率
    print(f" [Excel VM增长TOP100 第一台 VM] 名称: {top_sheet['B5'].value}, 分配空间: {top_sheet['E5'].value}, 当前容量: {top_sheet['F5'].value}, 使用率: {top_sheet['G5'].value}")

    # 3.4 日增长/月增长详情
    for s_name in ["日增长详情", "月增长详情"]:
        if s_name in wb.sheetnames:
            detail_sheet = wb[s_name]
            detail_headers = [detail_sheet.cell(row=3, column=col).value for col in range(1, 14)]
            print(f" [Excel {s_name}] 表头 13 列: {detail_headers}")
            assert "分配空间" in detail_headers and "使用率" in detail_headers, f"{s_name} 缺失分配空间或使用率！"
            print(f" [Excel {s_name} 第一台 VM] 名称: {detail_sheet['B4'].value}, 分配空间: {detail_sheet['G4'].value}, 当前容量: {detail_sheet['H4'].value}, 使用率: {detail_sheet['I4'].value}")

    # 3.5 新建 VM Sheet
    for sheet_name in ["本日新建VM", "本月新建VM"]:
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            headers_new = [sheet.cell(row=2, column=col).value for col in range(1, 10)]
            print(f" [Excel {sheet_name}] 表头 9 列: {headers_new}")
            assert "分配空间" in headers_new, f"{sheet_name} 缺失「分配空间」！"
            assert "当前容量" in headers_new, f"{sheet_name} 缺失「当前容量」！"
            assert "使用率" in headers_new, f"{sheet_name} 缺失「使用率」！"

    # 3.6 每集群 Sheet
    cluster_sheets = [s for s in wb.sheetnames if s not in ["封面", "执行摘要", "数据质量说明", "容量趋势", "VM增长TOP100", "日增长详情", "月增长详情", "本日新建VM", "本月新建VM"]]
    print(f" 发现集群 Sheet 数量: {len(cluster_sheets)}")
    for c_name in cluster_sheets:
        cs = wb[c_name]
        head2 = [cs.cell(row=2, column=col).value for col in range(1, 11)]
        print(f" [Excel 集群 {c_name}] 头部第 2 行: {head2}")
        assert "已分配容量" in head2, f"集群 {c_name} 头部缺失「已分配容量」！"
        # 查找 VM 表头
        vm_headers = None
        for r in range(5, cs.max_row + 1):
            if cs.cell(row=r, column=1).value == "VM":
                vm_headers = [cs.cell(row=r, column=col).value for col in range(1, 10)]
                break
        if vm_headers:
            print(f" [Excel 集群 {c_name}] VM 表头 9 列: {vm_headers}")
            assert "分配空间" in vm_headers, f"集群 {c_name} VM 表头缺失「分配空间」！"
            assert "使用率" in vm_headers, f"集群 {c_name} VM 表头缺失「使用率」！"

    # 4. 覆盖桌面文件
    desktop_xlsx = "/Users/nazawsze/Desktop/SmartX_存储容量分析报告_20261010.xlsx"
    desktop_docx = "/Users/nazawsze/Desktop/SmartX_存储容量分析报告_20261010.docx"
    with open(desktop_xlsx, "wb") as f:
        f.write(excel_bytes)
    with open(desktop_docx, "wb") as f:
        f.write(word_bytes)
    print(f"\n 已将最新 Excel 报表写入: {desktop_xlsx}")
    print(f" 已将最新 Word 报表写入: {desktop_docx}")

    print("\n 全量真实报表导出验证全部通过 (Word / Excel 虚拟机「分配空间」与「使用率」准确落位)！")

if __name__ == "__main__":
    main()
